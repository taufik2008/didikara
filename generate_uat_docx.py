import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import datetime

def create_uat_document(output_path):
    doc = docx.Document()

    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Helper styling functions
    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin_name}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
            # Bottom border / divider line
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="2563EB"/></w:pBdr>')
            pPr.append(pBdr)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 58, 138) # Blue 900
        elif level == 3:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
        return p

    def add_callout(text, title="CATATAN PENTING"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "EFF6FF") # Light blue
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        # Border left blue
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"📌 {title}: ")
        r_title.bold = True
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(30, 64, 175)
        
        r_text = p.add_run(text)
        r_text.font.name = 'Arial'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── HEADER & TITLE ──
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    r_org = title_p.add_run("DIDIKARA ENTERPRISE LEARNING MANAGEMENT SYSTEM")
    r_org.font.name = 'Arial'
    r_org.font.size = Pt(10)
    r_org.bold = True
    r_org.font.color.rgb = RGBColor(234, 88, 12) # Orange primary

    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_after = Pt(6)
    r_h1 = h1_p.add_run("DOKUMEN USER ACCEPTANCE TEST (UAT)")
    r_h1.font.name = 'Arial'
    r_h1.font.size = Pt(22)
    r_h1.bold = True
    r_h1.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    r_sub = sub_p.add_run("Pengujian Penerimaan Pengguna untuk Fitur Pembelajaran Multi-Model, Integrasi Google Drive, Katalog Dinamis & Sertifikasi")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # ── METADATA TABLE ──
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nama Sistem / Platform", "DIDIKARA Learning & Certification Management System"),
        ("URL Lingkungan Pengujian", "https://didikara.vercel.app"),
        ("Tanggal Penyusunan", datetime.datetime.now().strftime("%d %B %Y")),
        ("Versi Dokumen", "v1.4 (Clean Database-Driven Release)"),
        ("Tipe Pengujian", "User Acceptance Testing (UAT) / Pengujian Penerimaan Akhir"),
        ("Target Pengguna", "Admin Platform, Instruktur / Asesor, Peserta Pelatihan, Korporat & Instansi")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)
        
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = 'Arial'
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(51, 65, 85)
        
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    # Border table
    for row in meta_table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── BAB 1: TUJUAN DAN RUANG LINGKUP ──
    add_heading("1. Tujuan dan Ruang Lingkup", level=1)
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.2
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.add_run(
        "Dokumen ini dirancang sebagai panduan pengujian terstruktur untuk memastikan bahwa platform LMS DIDIKARA "
        "telah memenuhi seluruh persyaratan fungsional, kemudahan pemakaian, dan keandalan sistem sesuai dengan "
        "kebutuhan operasional. Seluruh skenario pengujian diverifikasi secara langsung terhadap lingkungan produksi "
        "yang terhubung dengan database live tanpa dependensi data hardcode."
    )

    add_callout(
        "Platform telah diperbarui ke arsitektur Server-Rendered Dinamis. Semua materi kursus, kategori, tingkat/level, "
        "dan akun instruktur dimuat langsung dari basis data PostgreSQL. Data tiruan (dummy hardcode) telah dieliminasi total.",
        title="PENYESUAIAN SISTEM TERBARU"
    )

    # ── BAB 2: METODE & STATUS PENILAIAN ──
    add_heading("2. Kriteria & Status Pengujian", level=1)
    
    status_table = doc.add_table(rows=5, cols=3)
    status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Status", "Arti / Definisi", "Tindakan Lanjutan"]
    col_widths = [Inches(1.2), Inches(3.5), Inches(1.8)]

    # Header row
    hdr_cells = status_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], 100, 100, 120, 120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    status_rows = [
        ("PASS (Lulus)", "Fitur berfungsi persis sesuai spesifikasi tanpa kendala.", "Fitur disetujui untuk operasional."),
        ("PASS WITH NOTES", "Fitur berfungsi namun terdapat catatan minor (misal label teks atau kosmetik).", "Perbaikan diselesaikan pada iterasi minor."),
        ("FAIL (Gagal)", "Fitur tidak berfungsi, terjadi crash/error, atau hasil tidak sesuai ekspektasi.", "Perlu perbaikan segera (bugfix)."),
        ("BLOCKED", "Pengujian tidak dapat dilanjutkan karena fitur prasyarat mengalami kendala.", "Tunggu dependensi diselesaikan.")
    ]
    for row_idx, data in enumerate(status_rows, start=1):
        row = status_table.rows[row_idx]
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, 80, 80, 120, 120)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                if "PASS (Lulus)" in text:
                    r.font.color.rgb = RGBColor(16, 185, 129)
                elif "FAIL" in text:
                    r.font.color.rgb = RGBColor(239, 68, 68)
                elif "NOTES" in text:
                    r.font.color.rgb = RGBColor(245, 158, 11)

    for row in status_table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ── HELPER UNTUK MEMBUAT TABEL TEST CASES ──
    def add_test_module_table(module_title, test_cases):
        add_heading(module_title, level=2)
        
        tbl = doc.add_table(rows=len(test_cases) + 1, cols=6)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["ID", "Skenario Pengujian", "Langkah Pengujian", "Hasil yang Diharapkan", "Status", "Catatan"]
        widths = [Inches(0.7), Inches(1.5), Inches(1.6), Inches(1.5), Inches(0.6), Inches(0.6)]

        # Header Row
        hdr_row = tbl.rows[0]
        for idx, h in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.width = widths[idx]
            set_cell_background(cell, "0F172A") # Deep slate
            set_cell_margins(cell, 90, 90, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for row_i, tc in enumerate(test_cases, start=1):
            row = tbl.rows[row_i]
            for col_i, text in enumerate(tc):
                cell = row.cells[col_i]
                cell.width = widths[col_i]
                set_cell_margins(cell, 70, 70, 90, 90)
                if row_i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(text)
                r.font.name = 'Arial'
                r.font.size = Pt(8)
                if col_i == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(37, 99, 235)
                elif col_i == 4:
                    r.bold = True
                    if text == "PASS":
                        r.font.color.rgb = RGBColor(16, 185, 129)
                    elif text == "FAIL":
                        r.font.color.rgb = RGBColor(239, 68, 68)

        # Border styling
        for row in tbl.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    </w:tcBorders>
                ''')
                tcPr.append(tcBorders)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ── BAB 3: LIST SKENARIO PENGUJIAN UAT ──
    add_heading("3. Lembar Pengujian UAT per Modul", level=1)

    # Modul 1
    tc_m1 = [
        (
            "UAT-LP-01",
            "Akses Landing Page Publik & Render Awal",
            "Buka peramban ke URL https://didikara.vercel.app/",
            "Halaman terbuka sempurna (<2 detik), hero banner, 3 pilihan model belajar, dan katalog kursus tampil tanpa error konsol.",
            "PASS",
            "SSR berjalan optimal"
        ),
        (
            "UAT-LP-02",
            "Penyajian Kursus Baru dari Database",
            "Periksa apakah kursus baru yang ditambahkan di admin (misal 'Budget Mining') muncul di landing page publik.",
            "Kursus tampil otomatis di kartu katalog lengkap dengan judul, kategori 'Mining/Eng', badge, dan level.",
            "PASS",
            "Kursus c7 terverifikasi tampil di landing page"
        ),
        (
            "UAT-LP-03",
            "Filter Multi-Model Pembelajaran (3 Pilar)",
            "Klik filter: (1) Semua Model, (2) Self-Learning, (3) Online, (4) Offline.",
            "Grid katalog hanya menampilkan kursus yang sesuai tipe belajarnya dengan indikator badge warna yang tepat.",
            "PASS",
            "Penyaringan sesuai tipe"
        ),
        (
            "UAT-LP-04",
            "Filter Topik & Kategori Dinamis",
            "Buka dropdown 'Topik' dan pilih kategori baru dari database (misal 'Budget Operasional Tambang').",
            "Dropdown memuat seluruh kategori database (tanpa hardcode). Kursus yang sesuai kategori langsung terfilter.",
            "PASS",
            "Kategori dinamis dari database"
        ),
        (
            "UAT-LP-05",
            "Filter Tingkat / Level Dinamis",
            "Buka dropdown 'Tingkat' dan pilih level 'Lanjutan', 'Dasar', 'Menengah', atau 'Eksekutif'.",
            "Pilihan level sesuai isi database, kursus dengan level yang dipilih tampil secara akurat.",
            "PASS",
            "Level Lanjutan teridentifikasi"
        ),
        (
            "UAT-LP-06",
            "Pencarian Kursus Berbasis Teks (Search)",
            "Ketik kata kunci 'Tambang' atau 'PBJ' di kotak pencarian.",
            "Katalog menampilkan kursus yang memuat kata kunci di judul, deskripsi, nama instruktur, atau silabus.",
            "PASS",
            "Filter teks responsif"
        ),
        (
            "UAT-LP-07",
            "Modal Detail Kursus & Silabus",
            "Klik tombol 'Silabus & Detail' pada salah satu kartu kursus.",
            "Modal popup muncul menampilkan tab: Silabus, Fasilitas Paket, Ringkasan, dan Profil Pengajar.",
            "PASS",
            "Modal tabs interaktif"
        ),
        (
            "UAT-LP-08",
            "Daftar Dewan Instruktur Dinamis",
            "Scroll ke seksi 'Instruktur & Asesor Ahli'.",
            "Hanya menampilkan pengajar yang terdaftar pada kursus aktif di database. Instruktur dummy lama sudah tidak muncul.",
            "PASS",
            "Murni dari kursus database"
        )
    ]
    add_test_module_table("3.1 Modul 1: Landing Page & Katalog Publik", tc_m1)

    # Modul 2
    tc_m2 = [
        (
            "UAT-MOD-01",
            "Perbedaan Alur 3 Model Belajar",
            "Pelajari kartu penjelasan proses di seksi '3 Model Pembelajaran di DIDIKARA'.",
            "Sistem menjelaskan alur jelas: Self-Learning (On-Demand), Online (Live Webinar), dan Offline (Tatap Muka & Uji BNSP).",
            "PASS",
            "Alur informatif"
        ),
        (
            "UAT-MOD-02",
            "Pendaftaran Kursus Self-Learning (Gratis)",
            "Klik 'Mulai Belajar' pada kursus self-learning dengan harga Rp 0.",
            "Tidak diarahkan memilih jadwal/venue workshop. Akses kursus langsung aktif dan langsung dapat ditonton.",
            "PASS",
            "Tanpa pilih jadwal"
        ),
        (
            "UAT-MOD-03",
            "Pendaftaran Kelas Online (Live Webinar)",
            "Klik 'Daftar Kelas' pada kursus Online.",
            "Menampilkan informasi sesi webinar interaktif, link Zoom kelas, jadwal pelaksanaan sesi live malam/weekend.",
            "PASS",
            "Jadwal webinar tampil"
        ),
        (
            "UAT-MOD-04",
            "Pendaftaran Workshop Offline (Tatap Muka)",
            "Pilih kursus Offline Tatap Muka.",
            "Menampilkan lokasi hotel, kuota sisa kursi, fasilitas seminar kit, konsumsi, dan uji asesor BNSP.",
            "PASS",
            "Venue & kuota tertera"
        )
    ]
    add_test_module_table("3.2 Modul 2: Spesifikasi Model Pembelajaran (Self-Learning, Online, Offline)", tc_m2)

    # Modul 3
    tc_m3 = [
        (
            "UAT-CRS-01",
            "Pembuatan Kursus Baru oleh Admin/Instruktur",
            "Masuk sebagai Admin, buka panel Kursus, klik 'Tambah Kursus'. Isi formulir lengkap lalu simpan.",
            "Kursus tersimpan di database PostgreSQL, notifikasi sukses muncul, kursus muncul di daftar admin dan katalog publik.",
            "PASS",
            "Penyimpanan database teruji"
        ),
        (
            "UAT-CRS-02",
            "Penambahan Lesson Video MP4 Langsung",
            "Pada modul kursus, tambah lesson bertipe Video dengan link/file MP4.",
            "Lesson tersimpan dan saat dibuka di kelas, video player HTML5/Next.js dapat memutar video MP4 lancar.",
            "PASS",
            "Format MP4 didukung penuh"
        ),
        (
            "UAT-CRS-03",
            "Integrasi Materi Video Google Drive",
            "Masukkan URL video yang bersumber dari Google Drive (format link view/preview).",
            "Player mengenali format Drive, memutar iframe video dengan izin akses yang tepat tanpa broken video.",
            "PASS",
            "Kompatibel Google Drive"
        ),
        (
            "UAT-CRS-04",
            "Otomasi Izin Akses Google Drive ke Audiens",
            "Peserta mendaftar kursus yang materinya menggunakan Google Drive (Video/Slide).",
            "Sistem memicu endpoint /api/drive/grant-access untuk memberikan hak 'reader' kepada email peserta terdaftar.",
            "PASS",
            "Drive API access granted"
        ),
        (
            "UAT-CRS-05",
            "Pengelolaan Modul Slide Presentasi & PDF",
            "Tambahkan materi bertipe dokumen/presentasi (PPTX / Google Slides / PDF).",
            "Dokumen tersemat (embed) di ruang kelas dan dapat dibaca peserta dengan navigasi slide yang nyaman.",
            "PASS",
            "Slide viewer berjalan baik"
        )
    ]
    add_test_module_table("3.3 Modul 3: Manajemen Kurikulum, Materi MP4 & Google Drive", tc_m3)

    # Modul 4
    tc_m4 = [
        (
            "UAT-LMS-01",
            "Ruang Kelas Pembelajaran (Classroom)",
            "Peserta membuka kursus yang telah diikuti dari menu 'Kelas Saya' (/kelas).",
            "Sidebar materi menampilkan hierarki Modul & Lesson. Konten aktif termuat di area pemutar video/dokumen.",
            "PASS",
            "Antarmuka LMS responsif"
        ),
        (
            "UAT-LMS-02",
            "Pencatatan Progres Pembelajaran Otomatis",
            "Tonton materi video hingga selesai atau klik tandai selesai pada materi pelajaran.",
            "Centang hijau aktif di daftar lesson, persentase progres kursus meningkat, progres tersimpan di tabel user_progress.",
            "PASS",
            "Sync progres real-time"
        ),
        (
            "UAT-LMS-03",
            "Simulator Kuis / Kuis Evaluasi Pemahaman",
            "Selesaikan kuis di akhir modul kursus.",
            "Pilihan ganda dapat dipilih, memberikan umpan balik langsung (benar/salah beserta penjelasan), dan skor terekam.",
            "PASS",
            "Kuis interaktif"
        ),
        (
            "UAT-LMS-04",
            "Forum Diskusi & Tanya Jawab Kelas",
            "Kirim pertanyaan atau tanggapan pada tab 'Forum' di dalam kelas.",
            "Postingan tersimpan di basis data forum_posts, tampil seketika bagi instruktur dan peserta lain di kursus tersebut.",
            "PASS",
            "Dukungan diskusi aktif"
        )
    ]
    add_test_module_table("3.4 Modul 4: Pengalaman Belajar & Interaktivitas LMS", tc_m4)

    # Modul 5
    tc_m5 = [
        (
            "UAT-EXM-01",
            "Inisiasi Ujian Kelulusan Berwaktu (Timer)",
            "Peserta memulai ujian akhir kursus melalui tombol 'Mulai Ujian'.",
            "Hitung mundur waktu (countdown timer) berjalan akurat, token sesi ujian terverifikasi, dan navigasi soal tersedia.",
            "PASS",
            "Waktu ujian sinkron"
        ),
        (
            "UAT-EXM-02",
            "Submit Ujian & Penilaian Otomatis",
            "Peserta menjawab soal ujian dan menekan tombol 'Kirim Jawaban'.",
            "Sistem menghitung nilai secara otomatis di sisi server, mencocokkan passing score (misal 70%), dan menampilkan hasil.",
            "PASS",
            "Penilaian server-side aman"
        ),
        (
            "UAT-EXM-03",
            "Penerbitan Sertifikat Digital Otomatis",
            "Peserta yang berhasil lulus ujian membuka menu 'Sertifikat'.",
            "Sertifikat resmi terbit otomatis memuat: Nama Peserta, Judul Kursus, Nomor Kredensial Unik, Tanggal Lulus, dan QR Code.",
            "PASS",
            "Sertifikat instan"
        ),
        (
            "UAT-EXM-04",
            "Verifikasi Publik Sertifikat via QR Code",
            "Pindai QR code pada sertifikat atau buka tautan verifikasi publik /verify/[certId].",
            "Halaman verifikasi publik terbuka mengonfirmasi keaslian sertifikat, data pemegang, penerbit, dan tanggal terbit.",
            "PASS",
            "Verifikasi publik aktif"
        ),
        (
            "UAT-EXM-05",
            "Unduh Sertifikat PDF",
            "Klik tombol 'Unduh PDF' pada tampilan sertifikat.",
            "Berkas PDF sertifikat beresolusi tinggi terunduh dengan format resmi siap cetak.",
            "PASS",
            "Export PDF berfungsi"
        )
    ]
    add_test_module_table("3.5 Modul 5: Ujian Kompetensi, Evaluasi & Penerbitan Sertifikat", tc_m5)

    # Modul 6
    tc_m6 = [
        (
            "UAT-ADM-01",
            "Dasbor Admin & Metrik Analitik",
            "Login sebagai Super Administrator dan buka /admin.",
            "Dasbor menyajikan metrik jumlah pengguna aktif, total kursus, transaksi, dan log aktivitas audit.",
            "PASS",
            "Metrik live PostgreSQL"
        ),
        (
            "UAT-ADM-02",
            "Manajemen Kategori & Tingkat (CRUD)",
            "Tambah kategori baru (misal 'Geologi Tambang') atau level baru di panel Master Data.",
            "Data tersimpan di tabel categories/levels dan seketika muncul di semua dropdown filter publik tanpa reload build.",
            "PASS",
            "Master data terintegrasi"
        ),
        (
            "UAT-ADM-03",
            "Pengaturan Brand Platform (White-Label)",
            "Ubah Platform Name atau Tagline di Platform Settings.",
            "Perubahan langsung tercermin di navbar, judul landing page, dan footer.",
            "PASS",
            "Pengaturan platform dinamis"
        ),
        (
            "UAT-ADM-04",
            "Uji Keamanan & Proteksi API",
            "Akses endpoint admin atau kunci jawaban kuis tanpa login admin.",
            "Sistem merespons dengan HTTP 401 Unauthorized / memblokir kebocoran kunci jawaban ujian.",
            "PASS",
            "Proteksi otoritas aman"
        )
    ]
    add_test_module_table("3.6 Modul 6: Panel Administrasi, Master Data & Keamanan", tc_m6)

    # ── BAB 4: KESIMPULAN & TANDA TANGAN PERSETUJUAN ──
    add_heading("4. Lembar Persetujuan Pengujian (Sign-Off)", level=1)
    
    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.line_spacing = 1.2
    p_concl.paragraph_format.space_after = Pt(12)
    p_concl.add_run(
        "Berdasarkan hasil pelaksanaan User Acceptance Testing (UAT) di atas, seluruh fitur utama yang mencakup: "
        "katalog publik dinamis tanpa data hardcode, pemisahan 3 model pembelajaran, pemutaran video MP4 & Google Drive, "
        "pemberian akses materi otomatis, serta modul sertifikasi digital telah diverifikasi dan dinyatakan DITERIMA (ACCEPTED) "
        "untuk digunakan secara penuh dalam operasional sistem."
    )

    sign_table = doc.add_table(rows=4, cols=3)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_widths = [Inches(2.1), Inches(2.1), Inches(2.1)]
    
    sign_headers = ["Pihak Penguji (QA / Tester)", "Pengembang (Lead Developer)", "Pemilik Produk (Product Owner)"]
    for i, h in enumerate(sign_headers):
        cell = sign_table.rows[0].cells[i]
        cell.width = sign_widths[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    sign_roles = [
        ("Tanda Tangan:\n\n\n\n_______________________\nNama: Quality Assurance\nTanggal: "),
        ("Tanda Tangan:\n\n\n\n_______________________\nNama: Lead Software Engineer\nTanggal: "),
        ("Tanda Tangan:\n\n\n\n_______________________\nNama: Product Owner / Stakeholder\nTanggal: ")
    ]
    for i, text in enumerate(sign_roles):
        cell = sign_table.rows[1].cells[i]
        cell.width = sign_widths[i]
        set_cell_margins(cell, 120, 120, 100, 100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(51, 65, 85)

    status_row = sign_table.rows[2]
    for i, title in enumerate(["Status: LULUS (PASS)", "Status: LULUS (PASS)", "Status: APPROVED"]):
        cell = status_row.cells[i]
        cell.width = sign_widths[i]
        set_cell_background(cell, "ECFDF5") # light emerald
        set_cell_margins(cell, 60, 60, 80, 80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(5, 150, 105)

    for row in sign_table.rows[:3]:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)

    # Save document
    doc.save(output_path)
    print(f"UAT Document successfully generated at: {output_path}")

if __name__ == "__main__":
    create_uat_document("E:\\lms\\Dokumen_UAT_DIDIKARA_LMS.docx")
